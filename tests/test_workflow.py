import json
from pathlib import Path
import subprocess
import sys
import shutil
import unittest
import uuid


ROOT = Path(__file__).resolve().parents[1]
CLI = ROOT / "ppt-drawer" / "scripts" / "ppt_drawer.py"


class WorkflowTests(unittest.TestCase):
    def setUp(self):
        temp_root = ROOT / "tmp" / "workflow-tests"
        temp_root.mkdir(parents=True, exist_ok=True)
        self.temp_root = temp_root.resolve()
        self.base = temp_root / ("PPT项目-" + uuid.uuid4().hex)
        self.base.mkdir()
        self.project = self.base / "项目一"
        self.project.mkdir()

    def tearDown(self):
        assert self.base.resolve().is_relative_to(self.temp_root)
        shutil.rmtree(self.base)

    def call(self, command, *args, ok=True, project=None):
        result = subprocess.run(
            [sys.executable, "-X", "utf8", str(CLI), command, "--project", str(project or self.project), "--run", "demo", *map(str, args)],
            capture_output=True, text=True, encoding="utf-8",
        )
        if ok:
            self.assertEqual(result.returncode, 0, result.stderr or result.stdout)
        else:
            self.assertNotEqual(result.returncode, 0)
        return json.loads(result.stdout)

    def test_project_materials_are_indexed_and_isolated(self):
        paper = self.project / "研究材料.md"
        paper.write_text("# Study objective\nMeasured signal: 12 units.\n", encoding="utf-8")
        self.call("init", "--title", "研究汇报")
        self.call("ingest", "--input", paper)
        result = self.call("status")
        self.assertEqual(result["title"], "研究汇报")
        self.assertEqual(len(result["materials"]), 1)
        self.assertEqual(result["materials"][0]["status"], "ready")
        self.assertIn("Measured signal: 12 units.", result["materials"][0]["content"][0]["text"])
        other = self.base / "项目二"
        other.mkdir()
        self.call("init", "--title", "另一份汇报", project=other)
        self.assertEqual(self.call("status", project=other)["materials"], [])

    def test_mixed_materials_report_locators_and_unreadable_files(self):
        from docx import Document
        from openpyxl import Workbook
        from PIL import Image
        from pypdf import PdfWriter
        from pypdf.generic import DecodedStreamObject, DictionaryObject, NameObject
        import zipfile
        materials = self.project / "材料"
        materials.mkdir()
        (materials / "notes.md").write_text("Measured sample A", encoding="utf-8")
        (materials / "data.csv").write_text("sample,value\nA,12\nB,18\n", encoding="utf-8")
        doc = Document()
        doc.add_paragraph("Document evidence")
        doc.save(materials / "paper.docx")
        wb = Workbook()
        wb.active.title = "实验"
        wb.active.append(["Sample", "Value"])
        wb.active.append(["A", 12])
        wb.active.append(["B", 0])
        wb.save(materials / "table.xlsx")
        Image.new("RGB", (120, 80), "white").save(materials / "figure.png")
        writer = PdfWriter()
        page = writer.add_blank_page(width=300, height=200)
        font = DictionaryObject({NameObject("/Type"):NameObject("/Font"), NameObject("/Subtype"):NameObject("/Type1"), NameObject("/BaseFont"):NameObject("/Helvetica")})
        page[NameObject("/Resources")] = DictionaryObject({NameObject("/Font"):DictionaryObject({NameObject("/F1"):writer._add_object(font)})})
        stream = DecodedStreamObject()
        stream.set_data(b"BT /F1 12 Tf 20 160 Td (PDF evidence) Tj ET")
        page[NameObject("/Contents")] = writer._add_object(stream)
        writer.write(materials / "paper.pdf")
        with zipfile.ZipFile(materials / "slides.pptx", "w") as z:
            z.writestr("ppt/presentation.xml", '<p:presentation xmlns:p="http://schemas.openxmlformats.org/presentationml/2006/main"><p:sldSz cx="12192000" cy="6858000"/></p:presentation>')
            z.writestr("ppt/slides/slide1.xml", '<p:sld xmlns:p="http://schemas.openxmlformats.org/presentationml/2006/main" xmlns:a="http://schemas.openxmlformats.org/drawingml/2006/main"><p:cSld><p:spTree><p:sp><p:txBody><a:p><a:r><a:t>PPT evidence</a:t></a:r></a:p></p:txBody></p:sp></p:spTree></p:cSld></p:sld>')
        (materials / "broken.pdf").write_bytes(b"not a pdf")
        own = self.project / ".ppt-assistant" / "ignored.md"
        own.parent.mkdir()
        own.write_text("Generated output must not be ingested", encoding="utf-8")
        self.call("init")
        self.call("ingest", "--input", self.project)
        result = self.call("status")
        records = {m["name"]:m for m in result["materials"]}
        self.assertEqual(set(records), {"notes.md","data.csv","paper.docx","table.xlsx","figure.png","paper.pdf","slides.pptx","broken.pdf"})
        self.assertEqual(records["broken.pdf"]["status"], "error")
        for name in ("notes.md","data.csv","paper.docx","table.xlsx","figure.png","paper.pdf","slides.pptx"):
            self.assertEqual(records[name]["status"], "ready", name)
            self.assertTrue(records[name]["sha256"], name)
            self.assertTrue(records[name]["content"][0]["locator"], name)
        self.assertIn("PDF evidence", records["paper.pdf"]["content"][0]["text"])
        self.assertEqual(records["table.xlsx"]["content"][0]["rows"][1], ["A",12])
        self.assertIn("B\t0",records["table.xlsx"]["content"][0]["text"])

    def write_json(self, name, value):
        path = self.project / name
        path.write_text(json.dumps(value, ensure_ascii=False), encoding="utf-8")
        return path

    def prepare_plan(self):
        material = self.project / "evidence.md"
        material.write_text("Measured signals: A=12, B=18", encoding="utf-8")
        self.call("init", "--title", "研究汇报")
        sid = self.call("ingest", "--input", material)["materials"][0]["id"]
        plan = {"version":1,"title":"研究汇报","style":"deep_mechanism","sample_slide_ids":["s1","s2","s3"],"slides":[{"id":f"s{i}","title":f"证据 {i}","layout":"title-content","takeaway":"Measured evidence","source_ids":[sid],"blocks":[{"type":"text","text":"A=12; B=18"}]} for i in range(1,5)]}
        path = self.write_json("plan-input.json", plan)
        self.call("plan", "--input", path)
        return sid, plan, path

    def test_knowledge_and_web_additions_require_traceable_sources(self):
        sid, _, _ = self.prepare_plan()
        good = self.write_json("knowledge.json", [{"id":"k1","kind":"project_result","text":"A measured signal is 12","source_ids":[sid]}])
        self.call("knowledge", "--input", good)
        bad = self.write_json("bad-knowledge.json", [{"id":"k2","kind":"project_result","text":"Untraceable","source_ids":["missing"]}])
        self.call("knowledge", "--input", bad, ok=False)
        additions = self.write_json("web.json", [{"id":"web_background","knowledge":"A background definition","url":"https://example.org/paper","purpose":"Explain the measurement","slide_ids":["s2"]}])
        self.call("supplement", "--input", additions)
        status = self.call("status")
        self.assertEqual(len(status["knowledge"]), 1)
        self.assertEqual(status["supplements"][0]["knowledge"], "A background definition")
        self.assertEqual(status["supplements"][0]["slide_ids"], ["s2"])
        report = Path(status["run_dir"]) / "supplemental-knowledge.md"
        self.assertIn("Explain the measurement", report.read_text(encoding="utf-8"))
        self.assertIn("https://example.org/paper", report.read_text(encoding="utf-8"))

    def test_plan_page_renaming_requires_explicit_supplement_mapping_update(self):
        _, plan, path = self.prepare_plan()
        addition = {"id":"web_background","knowledge":"A background definition","url":"https://example.org/paper","purpose":"Explain the measurement","slide_ids":["s2"]}
        additions = self.write_json("web.json", [addition])
        self.call("supplement", "--input", additions)
        self.call("register-sample", "--directory", self.sample_artifacts())
        approved = self.call("approve", "--statement", "测试确认：当前页面和来源")
        run = Path(approved["run_dir"])
        preserved = {name:(run/name).read_bytes() for name in ("plan.json","outline.md","supplemental-knowledge.md","state.json")}
        plan["slides"][1]["id"] = "new-page"
        plan["sample_slide_ids"][1] = "new-page"
        path.write_text(json.dumps(plan), encoding="utf-8")
        rejected = self.call("plan", "--input", path, ok=False)
        self.assertIn("web_background", rejected["error"])
        self.assertIn("slide_ids", rejected["error"])
        for name, content in preserved.items():
            self.assertEqual((run/name).read_bytes(), content, name)
        self.assertTrue(self.call("status")["approval_valid"])
        addition["slide_ids"] = []
        additions.write_text(json.dumps([addition]), encoding="utf-8")
        self.call("supplement", "--input", additions)
        self.call("plan", "--input", path)
        self.assertIn("尚未用于页面", (run/"supplemental-knowledge.md").read_text(encoding="utf-8"))
        addition["slide_ids"] = ["new-page"]
        additions.write_text(json.dumps([addition]), encoding="utf-8")
        result = self.call("supplement", "--input", additions)
        self.assertEqual(result["supplements"][0]["slide_ids"], ["new-page"])
        plan["slides"][1]["title"] = "Renamed title, same page identity"
        path.write_text(json.dumps(plan), encoding="utf-8")
        result = self.call("plan", "--input", path)
        self.assertEqual(result["supplements"][0]["slide_ids"], ["new-page"])

    def test_full_render_requires_approved_current_sample(self):
        self.prepare_plan()
        result = self.call("render", "--stage", "full", ok=False)
        self.assertIn("approval", result["error"].lower())

    def test_render_preflight_rejects_supplement_mapping_after_direct_plan_edit(self):
        _, plan, _ = self.prepare_plan()
        additions = self.write_json("web.json", [{"id":"web_background","knowledge":"A background definition","url":"https://example.org/paper","purpose":"Explain the measurement","slide_ids":["s2"]}])
        status = self.call("supplement", "--input", additions)
        plan["slides"][1]["id"] = "new-page"
        plan["sample_slide_ids"][1] = "new-page"
        (Path(status["run_dir"]) / "plan.json").write_text(json.dumps(plan), encoding="utf-8")
        rejected = self.call("render", "--stage", "sample", "--renderer", self.project/"missing-renderer.mjs", ok=False)
        self.assertIn("web_background", rejected["error"])
        self.assertIn("slide_ids", rejected["error"])
        self.assertIsNone(self.call("status")["sample"])

    def sample_artifacts(self, stage="sample"):
        import hashlib
        import zipfile
        from PIL import Image
        status = self.call("status")
        plan_path = Path(status["run_dir"]) / "plan.json"
        output = self.project / stage
        output.mkdir(exist_ok=True)
        with zipfile.ZipFile(output / "deck.pptx", "w") as archive:
            archive.writestr("ppt/presentation.xml", "<presentation/>")
            for i in range(1,4 if stage=="sample" else 5):
                archive.writestr(f"ppt/slides/slide{i}.xml", "<slide/>")
        for i in range(1,4 if stage=="sample" else 5):
            Image.new("RGB",(120,80),"white").save(output / f"slide-{i:03d}.png")
        (output / "render-report.json").write_text(json.dumps({"slide_ids":["s1","s2","s3"] if stage=="sample" else ["s1","s2","s3","s4"],"slide_count":3 if stage=="sample" else 4,"plan_sha256":hashlib.sha256(plan_path.read_bytes()).hexdigest(),"warnings":[]}),encoding="utf-8")
        return output

    def test_review_is_bound_to_plan_sources_and_sample_artifacts(self):
        _, plan, path = self.prepare_plan()
        output = self.sample_artifacts()
        self.call("register-sample", "--directory", output)
        before = self.call("status")
        self.assertEqual(before["stage"], "awaiting_review")
        self.assertFalse(before["approval_valid"])
        self.assertTrue(before["retrospectives"])
        self.call("approve", "--statement", "测试确认：已审阅当前大纲与三页样稿")
        self.assertTrue(self.call("status")["approval_valid"])
        (output / "deck.pptx").write_bytes((output / "deck.pptx").read_bytes()+b"changed")
        self.assertFalse(self.call("status")["approval_valid"])
        self.call("render", "--stage", "full", ok=False)
        self.call("register-sample", "--directory", output)
        self.call("approve", "--statement", "测试确认：复查更新样稿")
        plan["slides"][0]["takeaway"] = "Revised finding"
        path.write_text(json.dumps(plan),encoding="utf-8")
        self.call("plan", "--input", path)
        self.assertFalse(self.call("status")["approval_valid"])
        self.call("render", "--stage", "full", ok=False)

    def test_role_handoffs_and_verified_project_improvements_are_recorded(self):
        self.prepare_plan()
        output = self.project / "domain-output.json"
        output.write_text('{"facts":["A=12"]}',encoding="utf-8")
        self.call("handoff", "--role", "domain-analysis", "--agent-id", "/root/test-domain", "--input", self.project/"evidence.md", "--output", output, "--status", "completed", "--summary", "测试交接记录")
        findings = self.write_json("retro.json", {"summary":"验证后沉淀项目检查规则", "findings":[{"observation":"样本名应保持固定顺序","evidence":"provided dataset A then B","action":"增加项目检查规则","scope":"project","rule":"Keep samples in source order A then B","status":"verified","validation":"checked against data.csv rows"},{"observation":"通用布局建议尚待验证","evidence":"one draft page","action":"evaluate later","scope":"general","rule":"Unverified generic proposal","status":"proposed"}]})
        self.call("retro", "--stage", "sample", "--input", findings)
        status = self.call("status")
        self.assertEqual(status["handoffs"][0]["agent_id"], "/root/test-domain")
        self.assertEqual(status["handoffs"][0]["outputs"], [str(output.resolve())])
        self.assertEqual([p["rule"] for p in status["improvements"]], ["Keep samples in source order A then B"])
        self.assertEqual(status["retrospectives"][-1]["findings"][1]["status"], "proposed")

    def test_delivery_requires_current_full_artifacts_and_a_final_retrospective(self):
        self.prepare_plan()
        self.call("finish", ok=False)
        self.call("register-sample", "--directory", self.sample_artifacts())
        self.call("approve", "--statement", "测试确认：当前方案可继续")
        self.call("register-full", "--directory", self.sample_artifacts("full"))
        result = self.call("finish")
        self.assertEqual(result["stage"], "completed")
        self.assertEqual(result["full"]["slide_ids"], ["s1","s2","s3","s4"])
        self.assertEqual(result["retrospectives"][-1]["outcome"], "completed")
        self.assertTrue(result["approval_valid"])

    def test_changed_unreadable_source_cannot_be_reused_by_an_old_plan(self):
        self.prepare_plan()
        output = self.sample_artifacts()
        (self.project / "evidence.md").write_bytes(b"\xff\xfeinvalid source")
        result = self.call("ingest", "--input", self.project / "evidence.md")
        self.assertEqual(result["materials"][0]["status"], "error")
        rejected = self.call("register-sample", "--directory", output, ok=False)
        self.assertIn("source", rejected["error"].lower())

    def test_failed_render_preflight_leaves_a_failure_retrospective(self):
        self.prepare_plan()
        self.call("render", "--stage", "sample", "--renderer", self.project/"missing-renderer.mjs", ok=False)
        status = self.call("status")
        self.assertEqual(status["retrospectives"][-1]["outcome"], "failed")
        self.assertIn("adapter", status["last_error"].lower())


if __name__ == "__main__":
    unittest.main()
