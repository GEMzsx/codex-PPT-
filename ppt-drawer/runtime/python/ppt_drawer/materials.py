"""Read mixed project materials with explicit locators and completeness status."""
import csv
import hashlib
import io
import os
from pathlib import Path
import re
import xml.etree.ElementTree as ET
import zipfile

from .storage import file_hash


SUPPORTED = {".pdf", ".docx", ".md", ".txt", ".pptx", ".png", ".jpg", ".jpeg", ".webp", ".xlsx", ".csv"}
EXCLUDED = {".ppt-assistant", ".git", "node_modules", ".venv", "venv", "tmp", "__pycache__", ".pytest_cache", ".playwright-cli"}
MAX_TEXT = 250_000
MAX_ROWS = 1000
MAX_COLUMNS = 100
NS_A = "http://schemas.openxmlformats.org/drawingml/2006/main"


def discover(paths):
    found = {}
    for item in paths:
        path = Path(item).resolve()
        if path.is_dir():
            for parent, dirs, files in os.walk(path, followlinks=False):
                dirs[:] = sorted(d for d in dirs if d.lower() not in EXCLUDED and not (Path(parent) / d).is_symlink())
                for name in sorted(files):
                    candidate = Path(parent) / name
                    if candidate.suffix.lower() in SUPPORTED and not candidate.is_symlink():
                        found[str(candidate.resolve())] = candidate.resolve()
        else:
            found[str(path)] = path
    return list(found.values())


def _scalar(value):
    if value is None or isinstance(value, (str, int, float, bool)):
        return value
    return value.isoformat() if hasattr(value, "isoformat") else str(value)


def _table_rows(rows):
    values = []
    partial = False
    for i, row in enumerate(rows):
        if i >= MAX_ROWS:
            partial = True
            break
        row = list(row)
        if len(row) > MAX_COLUMNS:
            partial = True
        values.append([_scalar(v) for v in row[:MAX_COLUMNS]])
    return values, partial


def extract(path):
    ext = path.suffix.lower()
    content, warnings, assets = [], [], []
    if ext in {".md", ".txt"}:
        content = [{"locator":"document", "text":path.read_text(encoding="utf-8-sig")}]
    elif ext == ".csv":
        text = path.read_text(encoding="utf-8-sig")
        try:
            dialect = csv.Sniffer().sniff(text[:4096])
        except csv.Error:
            dialect = csv.excel
        rows, partial = _table_rows(csv.reader(io.StringIO(text), dialect))
        content = [{"locator":"rows 1-" + str(len(rows)), "rows":rows, "text":"\n".join("\t".join("" if x is None else str(x) for x in row) for row in rows)}]
        if partial:
            warnings.append("Table truncated at 1000 rows / 100 columns")
    elif ext == ".pdf":
        from pypdf import PdfReader
        reader = PdfReader(path)
        if reader.is_encrypted and not reader.decrypt(""):
            raise ValueError("Encrypted PDF requires a readable copy")
        for i, page in enumerate(reader.pages, 1):
            text = page.extract_text() or ""
            content.append({"locator":f"page {i}", "text":text})
            if not text.strip():
                warnings.append(f"page {i}: no extractable text; scanned/visual content needs inspection")
    elif ext == ".docx":
        from docx import Document
        document = Document(path)
        content = [{"locator":f"paragraph {i}", "text":p.text} for i, p in enumerate(document.paragraphs, 1) if p.text.strip()]
        for i, table in enumerate(document.tables, 1):
            rows, partial = _table_rows([[cell.text for cell in row.cells] for row in table.rows])
            content.append({"locator":f"table {i}", "rows":rows, "text":"\n".join("\t".join(row) for row in rows)})
            if partial:
                warnings.append(f"table {i}: truncated")
        with zipfile.ZipFile(path) as archive:
            assets = [{"locator":name, "kind":"embedded_image"} for name in archive.namelist() if name.startswith("word/media/") and not name.endswith("/")]
        if assets:
            warnings.append("Embedded images indexed; inspect visually before using their content")
    elif ext == ".pptx":
        with zipfile.ZipFile(path) as archive:
            slides = sorted((name for name in archive.namelist() if re.fullmatch(r"ppt/slides/slide\d+\.xml", name)), key=lambda n:int(re.search(r"slide(\d+)",n).group(1)))
            pres = ET.fromstring(archive.read("ppt/presentation.xml"))
            rel_name = "ppt/_rels/presentation.xml.rels"
            if rel_name in archive.namelist():
                import posixpath
                relations = {r.get("Id"):r.get("Target") for r in ET.fromstring(archive.read(rel_name))}
                order = []
                for item in pres.findall("{http://schemas.openxmlformats.org/presentationml/2006/main}sldIdLst/*"):
                    target = relations.get(item.get("{http://schemas.openxmlformats.org/officeDocument/2006/relationships}id"))
                    if target:
                        order.append(posixpath.normpath(posixpath.join("ppt",target)).lstrip("/"))
                if order:
                    slides = order
            for i, name in enumerate(slides, 1):
                root = ET.fromstring(archive.read(name))
                text = "\n".join(node.text or "" for node in root.iter("{" + NS_A + "}t"))
                content.append({"locator":f"slide {i}", "text":text, "part":name})
            assets = [{"locator":name,"kind":"embedded_media"} for name in archive.namelist() if name.startswith("ppt/media/") and not name.endswith("/")]
        if assets:
            warnings.append("Embedded media indexed; static text extraction does not inspect their visual content")
    elif ext == ".xlsx":
        from openpyxl import load_workbook
        workbook = load_workbook(path, read_only=True, data_only=False)
        try:
            for sheet in workbook:
                rows, partial = _table_rows(sheet.iter_rows(values_only=True))
                content.append({"locator":f"sheet {sheet.title}, rows 1-{len(rows)}", "sheet":sheet.title,"rows":rows,"text":"\n".join("\t".join("" if v is None else str(v) for v in row) for row in rows)})
                if partial:
                    warnings.append(f"sheet {sheet.title}: truncated at 1000 rows / 100 columns")
                if any(isinstance(v,str) and v.startswith("=") for row in rows for v in row):
                    warnings.append(f"sheet {sheet.title}: formulas retained; calculated values must be verified")
        finally:
            workbook.close()
    elif ext in {".png", ".jpg", ".jpeg", ".webp"}:
        from PIL import Image
        with Image.open(path) as image:
            image.verify()
        with Image.open(path) as image:
            content = [{"locator":"image", "text":"", "width":image.width,"height":image.height,"mode":image.mode}]
        assets = [{"locator":"image", "kind":"image", "path":str(path)}]
    else:
        raise ValueError("Unsupported format: " + ext)
    remaining = MAX_TEXT
    for part in content:
        text = part.get("text", "")
        if len(text) > remaining:
            part["text"] = text[:remaining]
            warnings.append("Extracted text truncated at 250000 characters")
        remaining = max(0, remaining - len(part.get("text", "")))
    return content, warnings, assets


def material_record(path):
    sid = "src_" + hashlib.sha256(os.path.normcase(str(path)).encode("utf-8")).hexdigest()[:12]
    result = {"id":sid,"path":str(path),"name":path.name,"format":path.suffix.lower(),"sha256":None,"status":"error","content":[],"warnings":[],"assets":[]}
    try:
        result["sha256"] = file_hash(path)
        content, warnings, assets = extract(path)
        result.update(content=content, warnings=warnings, assets=assets, status="partial" if warnings else "ready")
    except Exception as error:
        result["error"] = f"{type(error).__name__}: {error}"
    return result
